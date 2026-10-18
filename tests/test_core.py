"""Tes offline logika inti dengan LLM tiruan. Jalankan: python tests/test_core.py"""
import json
import os
import sys
import zipfile
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import core  # noqa: E402

PASSAGE = " ".join(["Kalimat wacana yang cukup panjang untuk pengujian."] * 20)
ARABIC = " ".join(["هذا نص عربي طويل للاختبار والتدريب على القراءة"] * 12)


def fake_batch(prompt_plan, lang=False, option_keys="ABCD", arabic=False):
    qs, bp, an = [], [], []
    for i, (num, t) in enumerate(prompt_plan):
        q = {"number": num, "type": t, "difficulty": "Sedang", "indicator": f"ind {num}",
             "passage_id": "P1" if lang else "", "stimulus": "",
             "question": f"Soal {num} $x^2$ <script>alert(1)</script>\nbaris kedua",
             "options": {}, "answer": "", "explanation": f"bahas {num}", "scoring_rubric": "rubrik"}
        if t == "Pilihan Ganda":
            q["options"] = {k: f"Opsi {k}" for k in option_keys}
            q["answer"] = "B. Opsi B"
        elif t == "Benar/Salah":
            q["options"] = {"Benar": "x", "Salah": "y"}
            q["answer"] = "true"
        else:
            q["answer"] = "jawaban"
        qs.append(q)
        bp.append({"number": num, "topic": "t", "competency": "c", "indicator": "i",
                   "cognitive_level": "C2", "difficulty": "Mudah", "question_type": "salah"})
        an.append({"number": num, "content_validity": "ok", "construction": "ok", "language": "ok", "notes": "-"})
    d = {"title": "Paket X", "instructions": "Kerjakan.", "questions": qs, "blueprint": bp,
         "qualitative_analysis": an, "material_summary": "ringkasan"}
    d["passages"] = [{"id": "P1", "title": "Judul", "text": ARABIC if arabic else PASSAGE}] if lang else []
    return d


def make_llm(cfg, calls, lang=False, flaky=False, arabic=False):
    plan_all = core.plan_types(cfg["qtypes"], cfg["count"])
    keys = "ABCDE" if cfg["option_format"] == "ABCDE" else "ABCD"

    def llm(messages, max_tokens, temperature=0.2):
        calls.append(max_tokens)
        prompt = messages[-1]["content"]
        if "Perbaiki teks berikut" in prompt:
            return json.dumps(fake_batch([(1, "Pilihan Ganda")], lang)), "stop"
        import re
        nums = [int(n) for n in re.findall(r"Soal nomor (\d+):", prompt)]
        plan = [(n, plan_all[n - 1]) for n in nums]
        raw = json.dumps(fake_batch(plan, lang, keys, arabic), ensure_ascii=False)
        if flaky and len(calls) == 1:
            return raw[: len(raw) // 2], "length"
        return "```json\n" + raw + "\n```", "stop"
    return llm


def base_cfg(**kw):
    cfg = {"subject": "Fisika", "grade": "10", "semester": "Semester 1", "topic": "t", "count": 7,
           "difficulty": "Campuran", "qtypes": ["Pilihan Ganda", "Essay", "Benar/Salah", "Isian Singkat"],
           "option_format": "ABCDE", "output_language": "Indonesia", "curriculum": "",
           "include_explanations": True, "extra_instructions": ""}
    cfg.update(kw)
    return cfg


def test_extract_json():
    assert core.extract_json('```json\n{"a": "x\ny",}\n``` selesai')["a"] == "x\ny"
    assert core.extract_json('<think>hm</think>{"a":1} trailing')["a"] == 1
    for bad in ("", "tidak ada json", '{"a": '):
        try:
            core.extract_json(bad)
        except ValueError:
            continue
        raise AssertionError(bad)


def test_generation_mixed_types():
    cfg, calls = base_cfg(), []
    data = core.generate_package(make_llm(cfg, calls), cfg)
    assert len(data["questions"]) == 7 and len(calls) == 3
    assert [q["number"] for q in data["questions"]] == list(range(1, 8))
    assert [b["number"] for b in data["blueprint"]] == list(range(1, 8))
    assert [q["type"] for q in data["questions"]] == core.plan_types(cfg["qtypes"], 7)
    assert all(b["question_type"] == q["type"] for b, q in zip(data["blueprint"], data["questions"]))
    pg = data["questions"][0]
    assert pg["answer"] == "B" and list(pg["options"]) == list("ABCDE")
    assert data["questions"][4]["answer"] == "Benar"  # urutan: PG, PG, Essay, Essay, B/S, B/S, Isian
    assert "\\n" not in data["material_summary"] and data["material_summary"].count("\n\n") == 2


def test_truncation_retry_raises_budget():
    cfg, calls = base_cfg(count=3), []
    core.generate_package(make_llm(cfg, calls, flaky=True), cfg)
    assert len(calls) == 2 and calls[1] > calls[0]


def test_option_count_mismatch_fails_cleanly():
    cfg = base_cfg(count=3, qtypes=["Pilihan Ganda"], option_format="ABCDE")
    plan = core.plan_types(cfg["qtypes"], 3)
    bad = fake_batch(list(enumerate(plan, 1)), option_keys="ABCD")
    try:
        core.normalize_batch(bad, plan, "ABCDE", cfg["qtypes"], "Fisika", 1)
    except ValueError as e:
        assert "opsi" in str(e).lower()
    else:
        raise AssertionError("harus gagal")


def test_language_passage_and_exports():
    cfg = base_cfg(subject="Bahasa Arab", output_language="Arab", count=5, qtypes=["Pilihan Ganda", "Essay"])
    data = core.generate_package(make_llm(cfg, [], lang=True, arabic=True), cfg)
    assert len(data["passages"]) == 2
    layout = core.passage_layout(data)
    assert sorted(layout) == [1, 4]
    meta = {"subject": "Bahasa Arab", "grade": "10", "semester": "Semester 1", "topic": "t",
            "output_language": "Arab", "difficulty": "Campuran", "qtype": "Pilihan Ganda, Essay"}
    blob = core.build_docx(data, meta)
    with zipfile.ZipFile(BytesIO(blob)) as z:
        xml = z.read("word/document.xml").decode()
    assert "<w:bidi/>" in xml and "<w:rtl/>" in xml
    assert xml.count('w:val="clear"') >= 5
    from docx import Document
    assert len(Document(BytesIO(blob)).tables) >= 2
    html = core.build_preview_html(data)
    assert "<script>" not in html and "\n" not in html and "$" not in html
    assert "Wacana untuk soal nomor" in html


def test_missing_passage_rejected():
    cfg = base_cfg(subject="Bahasa Indonesia", count=3, qtypes=["Pilihan Ganda"], option_format="ABCD")
    plan = core.plan_types(cfg["qtypes"], 3)
    d = fake_batch(list(enumerate(plan, 1)), lang=False)
    try:
        core.normalize_batch(d, plan, "ABCD", cfg["qtypes"], "Bahasa Indonesia", 1)
    except ValueError as e:
        assert "wacana" in str(e).lower()
    else:
        raise AssertionError("harus gagal")


def test_plan_types_grouped_per_type():
    plan = core.plan_types(["Essay", "Pilihan Ganda"], 20)  # urutan klik tidak berpengaruh
    assert plan == ["Pilihan Ganda"] * 10 + ["Essay"] * 10
    assert core.plan_types(["Pilihan Ganda", "Essay"], 7) == ["Pilihan Ganda"] * 4 + ["Essay"] * 3
    assert core.plan_types([], 3) == ["Pilihan Ganda"] * 3


def test_merge_regroups_when_ai_deviates_from_plan():
    cfg = base_cfg(count=6, qtypes=["Pilihan Ganda", "Essay"])
    types = ["Essay", "Pilihan Ganda", "Essay", "Pilihan Ganda", "Pilihan Ganda", "Essay"]
    plan = list(types)
    b1 = core.normalize_batch(fake_batch(list(enumerate(plan[:3], 1)), option_keys="ABCDE"),
                              plan[:3], "ABCDE", cfg["qtypes"], "Fisika", 1)
    b2 = core.normalize_batch(fake_batch(list(enumerate(plan[3:], 4)), option_keys="ABCDE"),
                              plan[3:], "ABCDE", cfg["qtypes"], "Fisika", 2)
    data = core.merge_batches([b1, b2], cfg)
    assert [q["type"] for q in data["questions"]] == ["Pilihan Ganda"] * 3 + ["Essay"] * 3
    assert [q["number"] for q in data["questions"]] == list(range(1, 7))
    assert [b["question_type"] for b in data["blueprint"]] == [q["type"] for q in data["questions"]]
    assert [a["number"] for a in data["qualitative_analysis"]] == list(range(1, 7))


def test_docx_sections_pg_then_essay():
    from docx import Document
    cfg, calls = base_cfg(count=20, qtypes=["Pilihan Ganda", "Essay"]), []
    data = core.generate_package(make_llm(cfg, calls), cfg)
    assert [q["type"] for q in data["questions"]] == ["Pilihan Ganda"] * 10 + ["Essay"] * 10
    secs = core.type_sections(data["questions"])
    assert [s["label"] for s in secs] == ["A. PILIHAN GANDA (Soal No. 1–10)", "B. ESSAY (Soal No. 11–20)"]
    meta = {"subject": "Fisika", "grade": "10", "semester": "Semester 1", "topic": "t",
            "output_language": "Indonesia", "difficulty": "Campuran", "qtype": "Pilihan Ganda, Essay"}
    doc = Document(BytesIO(core.build_docx(data, meta)))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    # soal, kunci, pembahasan, rubrik, kisi-kisi, analisis -> masing-masing A lalu B
    assert headings.count("A. PILIHAN GANDA (Soal No. 1–10)") == 6
    assert headings.count("B. ESSAY (Soal No. 11–20)") == 6
    assert "Jumlah soal: 20 (Pilihan Ganda: 10, Essay: 10)" in doc.tables[0].cell(0, 1).text
    q_pg, q_es = doc.tables[1], doc.tables[2]
    assert [r.cells[0].text for r in q_pg.rows] == [str(n) for n in range(1, 11)]
    assert [r.cells[0].text for r in q_es.rows] == [str(n) for n in range(11, 21)]
    assert all("A. " in r.cells[1].text for r in q_pg.rows)
    assert all("Jawaban:" in r.cells[1].text for r in q_es.rows)
    key_grid = doc.tables[3]  # kunci PG: kisi ringkas No./Jawaban
    assert [c.text for c in key_grid.rows[0].cells] == ["No."] + [str(n) for n in range(1, 11)]
    assert [c.text for c in key_grid.rows[1].cells] == ["Jawaban"] + ["B"] * 10
    html = core.build_preview_html(data)
    assert html.index("A. PILIHAN GANDA") < html.index("B. ESSAY")


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_"):
            fn()
            print("OK ", name)
